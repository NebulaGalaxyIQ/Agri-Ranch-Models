"""Source-backed V10 report engine.

Only allowlisted tenant-scoped ORM models are accessed. Report totals and exports
are built from the identical filtered query. No cross-currency addition, fake
observations, blind SQL fragments or unbounded client-side downloads.
"""
from __future__ import annotations
from collections import OrderedDict, defaultdict
from datetime import date, datetime, time, timedelta, timezone
from decimal import Decimal
from io import BytesIO, StringIO
import csv, math, re
from sqlalchemy import func, or_, String as SQLString, cast
from sqlalchemy.orm import Session
from sqlalchemy.sql.sqltypes import Date as SQLDate, DateTime as SQLDateTime, Numeric as SQLNumeric, Float as SQLFloat, Integer as SQLInteger
from fastapi import HTTPException
from database import Base
# Ensure ORM models register for both app and standalone test invocations.
from modules.crops import models as _crops
from modules.fattening import models as _fattening
from modules.livestock import models as _livestock
from modules.inventory import models as _inventory
from modules.finance import models as _finance
from modules.machinery import models as _machinery
from modules.employees import models as _employees
from modules.gardening import models as _gardening
from modules.organizations import models as _organizations
from modules.reports import v10_models as _report_models
from modules.organizations.service import permitted
from modules.reports.v10_analysis import label, analyses

# Module, display name, required read permission, primary dataset tables.
# Sources deliberately exclude user credentials, recovery keys and private account data.
DEFINITIONS=OrderedDict([
('executive',('Executive Dashboard','Executive','analytics.read',['v5_projects','v61_livestock_projects','v9_crop_projects','v5_finance_entries','v5_animals','v9_crop_harvest_lots','v6_machines'])),
('portfolio',('Project Portfolio','Executive','analytics.read',['v5_projects','v61_livestock_projects','v9_crop_projects','v9_crop_budgets','v61_livestock_project_costs'])),
('ranches',('Ranches & Infrastructure','Executive','ranch.manage',['v5_ranches','v8_livestock_houses','v9_crop_units','v6_crop_fields'])),
('fattening_population',('Fattening Operations & Population','Animal Fattening','animals.read',['v5_projects','v5_batches','v5_pens','v5_animals'])),
('fattening_weights',('Animal Weighing & Growth','Animal Fattening','animals.read',['v5_animals','v5_weight_observations'])),
('fattening_feed',('Feeding & Feed Management','Animal Fattening','rations.read',['v5_feeding_lines','v5_feed_plans','v5_rations','v5_mix_production','v5_routine_entries'])),
('fattening_sales',('Animal Sales & Break-even','Animal Fattening','finance.read',['v5_animal_sales','fattening_sale_cost_snapshots','v5_finance_entries'])),
('livestock_groups',('Herds, Flocks & Housing','Livestock & Poultry','livestock.read',['v6_livestock_groups','v83_livestock_group_profiles','v8_livestock_houses','v83_livestock_housing_profiles','v8_livestock_placements'])),
('livestock_registration',('Animals & Birds Registration','Livestock & Poultry','livestock.read',['v6_general_animals','v83_livestock_individual_profiles'])),
('livestock_breeding',('Breeding, Pregnancy & Offspring','Livestock & Poultry','livestock.read',['v8_livestock_breeding_records','v83_livestock_breeding_timeline','v83_livestock_offspring_cohorts'])),
('livestock_hatchery',('Poultry, Eggs & Incubation','Livestock & Poultry','livestock.read',['v8_livestock_incubations','v83_livestock_incubation_sources','v61_livestock_product_lots','v6_livestock_events'])),
('livestock_health',('Veterinary Care & Growth','Livestock & Poultry','livestock.read',['v8_livestock_health_records','v8_livestock_weight_records','v8_livestock_care_records'])),
('livestock_products',('Livestock Product Stock & Sales','Livestock & Poultry','livestock.read',['v61_livestock_product_lots','v61_livestock_product_sales','v6_livestock_sales'])),
('crop_planting',('Crop Projects, Fields & Planting','Crop Farming','crops.read',['v9_crop_projects','v9_crop_units','v9_crop_varieties','v9_crop_cycles','v6_crop_fields','v6_crop_cycles'])),
('crop_fieldwork',('Soil, Irrigation, Nutrition & Protection','Crop Farming','crops.read',['v9_crop_records','v6_crop_records','v61_crop_inputs','v9_crop_routines','v9_crop_completions'])),
('crop_trials',('Crop Trials & Experiments','Crop Farming','crops.read',['v9_crop_trials','v9_crop_trial_plots','v9_crop_records'])),
('crop_harvest',('Crop Harvest, Storage & Sales','Crop Farming','crops.read',['v9_crop_harvest_lots','v9_crop_stock_movements','v9_crop_sales','v6_crop_sales'])),
('crop_finance',('Crop Budgets, Costs & Outcomes','Crop Farming','crops.read',['v9_crop_budgets','v9_crop_projects','v9_crop_records','v9_crop_harvest_lots','v9_crop_sales'])),
('inventory',('Inventory & Procurement','Operations','stock.read',['v5_stock_items','v5_stock_movements','v57_stock_lots','v57_lot_issues','v5_stock_receipts'])),
('machinery',('Machinery & Assets','Operations','machinery.read',['v6_machines','v6_machine_records'])),
('employees',('Staff, Tasks & Routines','Operations','tasks.manage',['v5_tasks','v5_task_completions','v83_livestock_routine_completions','v9_crop_completions'])),
('gardening',('Gardening & Grounds','Operations','gardening.read',['v6_garden_areas','v6_garden_records'])),
('finance',('Finance & Financial Statements','Finance & Sales','finance.read',['v5_finance_entries','v9_crop_budgets','v61_livestock_project_costs','v9_crop_sales'])),
('commercial',('Sales, Buyers & Commercial Performance','Finance & Sales','finance.read',['v5_animal_sales','v61_livestock_product_sales','v9_crop_sales','v6_livestock_sales','v6_crop_sales'])),
('compliance',('Compliance & Audit','Governance','analytics.read',['v9_crop_records','v8_livestock_health_records','v5_tasks'])),
('custom',('Custom Section Explorer','Custom Reports','analytics.read',[])),
])
# Detailed operational report definitions. Internal storage names never appear in report titles.
MORE_REPORTS=OrderedDict([
 ('individual_animals',('Individual Animal Growth & History','Animal Fattening','animals.read',['v5_animals','v5_weight_observations','v5_animal_sales','fattening_sale_cost_snapshots','v5_finance_entries'])),
 ('fattening_acquisition',('Animal Purchases, Origins & Entry Weights','Animal Fattening','animals.read',['v5_animals','v5_batches','v5_projects'])),
 ('feed_consumption',('Actual Feed Consumption by Animal & Batch','Animal Fattening','rations.read',['v5_feeding_lines','v5_routine_entries','v5_stock_movements','v5_feed_plans'])),
 ('feed_formulations',('Feed Formulas, Ingredients & Production','Animal Fattening','rations.read',['v5_rations','v61_mix_reference_amounts','v5_mix_production','v5_feed_nutrition'])),
 ('fattening_project_costs',('Fattening Projects, Sales & Financial Results','Animal Fattening','finance.read',['v5_projects','v5_batches','v5_animals','v5_finance_entries','v5_animal_sales','fattening_sale_cost_snapshots'])),
 ('animal_sales_detail',('Individual Animal Sales & Cost Breakdown','Animal Fattening','finance.read',['v5_animal_sales','fattening_sale_cost_snapshots','v5_animals','v5_finance_entries'])),
 ('livestock_housing',('Housing Utilization & Animal Assignments','Livestock & Poultry','livestock.read',['v8_livestock_houses','v8_livestock_placements','v8_livestock_housing_movements','v83_livestock_housing_profiles'])),
 ('livestock_offspring',('Animal Births, Hatchings & Pedigree','Livestock & Poultry','livestock.read',['v83_livestock_offspring_cohorts','v83_livestock_individual_profiles','v8_livestock_breeding_records'])),
 ('livestock_costs',('Livestock Project Expenses & Allocations','Livestock & Poultry','livestock.read',['v61_livestock_projects','v61_livestock_project_costs','v83_livestock_registration_costs','v8_livestock_care_records','v8_livestock_health_records'])),
 ('livestock_routines',('Animal Daily Routines & Completion','Livestock & Poultry','livestock.read',['v83_livestock_routine_templates','v83_livestock_routine_completions','v8_livestock_care_records'])),
 ('livestock_production',('Milk, Egg, Wool & Other Production','Livestock & Poultry','livestock.read',['v6_livestock_events','v61_livestock_product_lots','v61_livestock_product_sales'])),
 ('crop_soils',('Field Soil Results & Land Preparation','Crop Farming','crops.read',['v6_crop_fields','v9_crop_units','v9_crop_records','v6_crop_records'])),
 ('crop_irrigation',('Water, Weather & Irrigation Events','Crop Farming','crops.read',['v9_crop_records','v9_crop_routines','v9_crop_completions'])),
 ('crop_nutrition',('Fertilizer, Crop Nutrition & Inputs','Crop Farming','crops.read',['v9_crop_records','v61_crop_inputs'])),
 ('crop_protection',('Pest Scouting, Disease & Treatments','Crop Farming','crops.read',['v9_crop_records','v61_crop_inputs'])),
 ('crop_growth',('Plant Growth, Monitoring & Forecasts','Crop Farming','crops.read',['v9_crop_records','v9_crop_cycles'])),
 ('crop_budget_detail',('Crop Budgets, Actual Spend & Variance','Crop Farming','crops.read',['v9_crop_budgets','v9_crop_records','v9_crop_sales','v9_crop_projects'])),
 ('crop_trial_detail',('Trial Treatments, Plots & Observations','Crop Farming','crops.read',['v9_crop_trials','v9_crop_trial_plots','v9_crop_records'])),
 ('crop_produce_stock',('Crop Harvest Lots, Losses & Available Stock','Crop Farming','crops.read',['v9_crop_harvest_lots','v9_crop_stock_movements','v9_crop_sales'])),
 ('crop_sales_detail',('Detailed Crop Sales & Product Yield','Crop Farming','crops.read',['v9_crop_sales','v9_crop_harvest_lots','v9_crop_cycles','v9_crop_varieties'])),
 ('stock_balances',('Stock Receipts, Issues & Current Balances','Operations','stock.read',['v5_stock_items','v5_stock_movements','v57_stock_lots','v57_lot_issues'])),
 ('stock_procurement',('Purchasing, Suppliers & Stock Receipts','Operations','stock.read',['v5_stock_receipts','v5_stock_movements','v5_stock_items'])),
 ('equipment_maintenance',('Machinery Operations, Repairs & Expenses','Operations','machinery.read',['v6_machines','v6_machine_records'])),
 ('staff_activity',('Staff Tasks, Shifts & Daily Completion','Operations','tasks.manage',['v5_tasks','v5_task_completions','v83_livestock_routine_completions','v9_crop_completions'])),
 ('finance_transactions',('Complete Recorded Income & Expenditure','Finance & Sales','finance.read',['v5_finance_entries'])),
 ('operating_project_results',('Production Project Financial Position','Finance & Sales','finance.read',['v5_finance_entries','v5_batches','v5_projects','v61_livestock_project_costs','v9_crop_budgets','v9_crop_records','v9_crop_sales'])),
])
DEFINITIONS.update(MORE_REPORTS)
# Any individual tenant-scoped physical table may be retrieved by the Section Explorer,
# subject to its owning module's permission and denylist.
EXPLORER_TABLES={source for spec in DEFINITIONS.values() for source in spec[3]} | {'v9_crop_routines','v9_crop_completions','v83_livestock_routine_templates','v83_livestock_routine_completions','v5_routine_entries','v5_stock_movements','v6_machine_records'}
HIDDEN_TABLES={'users','user_sessions','audit_logs','v57_recovery_codes','v57_org_roles','v57_member_roles',
 'v10_saved_reports','v10_report_runs','v10_report_schedules','v5_notifications','v5_organizations',
 'v5_organization_members','v5_memberships','v5_subscriptions','v5_license_keys','v5_feature_usage',
 'v5_audit_logs','v5_routine_period_claims'}
HIDDEN_FIELDS={'password_hash','token_hash','secret','raw_payload','access_token','api_key','password','recovery_code'}
DOMAIN_PREFIX=[('v9_crop','crops.read'),('v6_crop','crops.read'),('v61_crop','crops.read'),
 ('v83_livestock','livestock.read'),('v8_livestock','livestock.read'),('v6_livestock','livestock.read'),('v61_livestock','livestock.read'),
 ('v6_general_animals','livestock.read'),('v5_finance','finance.read'),('v5_stock','stock.read'),('v57_','stock.read'),('v6_machine','machinery.read'),('v6_garden','gardening.read'),
 ('v5_tasks','tasks.manage'),('v5_task_completions','tasks.manage'),('v5_feed','rations.read'),('v5_rations','rations.read'),('v5_mix','rations.read')]
DATE_CANDIDATES=['occurred_on','event_date','sold_on','sale_date','harvested_on','weigh_date','weighed_on','service_date','planting_date','collected_on','completed_at','recorded_on','received_on','issued_at','applied_on','inspection_date','delivery_date','entry_date','arrival_date','birth_date','start_date','due_date','day','date','recorded_at','created_at']
GROUP_CANDIDATES=['status','kind','category','species','crop','crop_name','product','name','purpose','grade','shift','action','type','currency','treatment','method']
VALUE_CANDIDATES=['amount','gross','revenue','total','quantity_kg','quantity','marketable_kg','gross_kg','weight_kg','current_weight_kg','head_count','area_ha','cost','price_per_kg','kg']
SCOPE_FIELDS=('project_id','ranch_id','batch_id','group_id','cycle_id','animal_id','item_id','machine_id','trial_id','lot_id','field_id','house_id','species','status','currency')
# A specialist report must not silently mix unrelated generic crop activities.
CROP_RECORD_SCOPES={
 'crop_soils':'soil','crop_irrigation':'irrigation','crop_nutrition':'nutrition',
 'crop_protection':'protection','crop_growth':'growth','crop_trial_detail':'trial_observations',
 'crop_trials':'trial_observations','crop_budget_detail':'costs',
}
def restrict_section(query,model,section_id):
    # Feed and other operational finance entries from different farm modules
    # must never appear as fattening expenses merely because a date matches.
    if (section_id in {'fattening_project_costs','fattening_sales','fattening_population',
                       'fattening_weights','individual_animals','animal_sales_detail'}
        and model.__tablename__=='v5_finance_entries'):
        query=query.filter(model.batch_id.isnot(None))
    target=CROP_RECORD_SCOPES.get(section_id)
    if target and model.__tablename__=='v9_crop_records':
        return query.filter(model.section==target)
    return query
MAX_EXPORT_ROWS=50000

def models_by_table():
    return {m.local_table.name:m.class_ for m in Base.registry.mappers if hasattr(m.class_,'__table__')}

def table_right(table):
    if table in ('v5_animal_sales','fattening_sale_cost_snapshots'):return 'finance.read'
    if table.startswith('v61_mix_'):return 'rations.read'
    for prefix,right in DOMAIN_PREFIX:
        if table.startswith(prefix):return right
    if table.startswith(('v5_animal','fattening_','v5_weight','v5_pens','v5_batches','v5_projects')):return 'animals.read'
    if table.startswith('v5_ranches'):return 'ranch.manage'
    if table.startswith('v5_routine'):return 'routine.read'
    return 'analytics.read'

def allowed(ctx,perm):
    member=ctx[2]
    try: permitted(member,perm);return True
    except HTTPException:return False

def all_tables(ctx):
    result=[]
    for table,model in sorted(models_by_table().items()):
        if table in HIDDEN_TABLES or table not in EXPLORER_TABLES or 'organization_id' not in model.__table__.columns:continue
        if not allowed(ctx,table_right(table)):continue
        result.append(table)
    return result

def report_sources(section,ctx):
    if section=='custom':return []
    definition=DEFINITIONS.get(section)
    if definition is None:
        if section.startswith('table:') and section[6:] in all_tables(ctx):return [section[6:]]
        raise HTTPException(404,'Unknown or unavailable report')
    if not allowed(ctx,definition[2]):raise HTTPException(403,'You do not have permission to view this report')
    # Do not silently include sources from modules the current role cannot read.
    return [t for t in definition[3] if t in all_tables(ctx)]

def catalogue(ctx):
    # Older report identifiers remain resolvable for saved report views, while the
    # navigation offers one clear entry per Fattening task, not repeated variants.
    hidden_repeats={'fattening_weights','fattening_sales','fattening_acquisition',
                    'feed_consumption','feed_formulations','animal_sales_detail'}
    arr=[dict(id=k,title=v[0],group=v[1],source_count=len(report_sources(k,ctx))) for k,v in DEFINITIONS.items() if allowed(ctx,v[2]) and k not in hidden_repeats]
    extra=[dict(id='table:'+t,title=label(t),group='Individual Section',source_count=1) for t in all_tables(ctx)]
    return {'sections':arr,'explorer':extra,'modules':list(dict.fromkeys(x['group'] for x in arr)),
            'note':'Reports are source-backed, scoped to your organization and module permissions. Financial amounts are never summed across currencies.'}

def columns(model):
    return [c.name for c in model.__table__.columns if c.name not in HIDDEN_FIELDS and not any(h in c.name for h in ('password','secret','token_hash'))]

def date_column(model, requested=None):
    cols=model.__table__.columns
    if requested:
        if requested not in cols or not isinstance(cols[requested].type,(SQLDate,SQLDateTime)):
            raise HTTPException(422,f'Date field {requested} is not supported by {model.__tablename__}')
        return requested
    for c in DATE_CANDIDATES:
        if c in cols and isinstance(cols[c].type,(SQLDate,SQLDateTime)):return c
    return None

def inherited_filter(db,model,organization_id,key,value):
    """Constrain indirect project/ranch relationships without copying parent data.
    No unverified relationship means no rows, not unfiltered output.
    """
    tab=model.__tablename__;c=model.__table__.columns
    if key not in ('project_id','ranch_id'):return None
    if 'batch_id' in c and tab.startswith(('v5_','fattening_')):
        from modules.organizations.models import V5Batch
        parent=db.query(V5Batch.id).filter(V5Batch.organization_id==organization_id,getattr(V5Batch,key)==value)
        return c.batch_id.in_(parent)
    if 'animal_id' in c and tab.startswith(('v5_','fattening_')):
        from modules.fattening.models import V5Animal
        from modules.organizations.models import V5Batch
        batches=db.query(V5Batch.id).filter(V5Batch.organization_id==organization_id,getattr(V5Batch,key)==value)
        animals=db.query(V5Animal.id).filter(V5Animal.organization_id==organization_id,V5Animal.batch_id.in_(batches))
        return c.animal_id.in_(animals)
    if key=='project_id' and 'group_id' in c and tab.startswith(('v6_livestock','v8_livestock','v83_livestock','v61_livestock')):
        from modules.livestock.models import LivestockProjectGroup
        parents=db.query(LivestockProjectGroup.group_id).filter(LivestockProjectGroup.organization_id==organization_id,LivestockProjectGroup.project_id==value)
        return c.group_id.in_(parents)
    if key=='project_id' and 'trial_id' in c and tab.startswith('v9_crop'):
        from modules.crops.models import V9Trial
        parents=db.query(V9Trial.id).filter(V9Trial.organization_id==organization_id,V9Trial.project_id==value)
        return c.trial_id.in_(parents)
    if key=='project_id' and 'lot_id' in c and tab.startswith('v9_crop'):
        from modules.crops.models import V9Harvest
        parents=db.query(V9Harvest.id).filter(V9Harvest.organization_id==organization_id,V9Harvest.project_id==value)
        return c.lot_id.in_(parents)
    return None

def q_for(db,model,organization_id, filters):
    c=model.__table__.columns
    q=db.query(model).filter(c.organization_id==organization_id)
    skipped=[]
    for key in SCOPE_FIELDS:
        v=getattr(filters,key,None)
        if v is not None:
            if key not in c:
                root_id_tables={'project_id':{'v5_projects','v9_crop_projects','v61_livestock_projects'},
                    'animal_id':{'v5_animals','v6_general_animals'},
                    'item_id':{'v5_stock_items'},'machine_id':{'v6_machines'},
                    'trial_id':{'v9_crop_trials'},'lot_id':{'v9_crop_harvest_lots','v61_livestock_product_lots'},
                    'field_id':{'v6_crop_fields'},'house_id':{'v8_livestock_houses'}}
                if key in root_id_tables and model.__tablename__ in root_id_tables[key]:
                    q=q.filter(c.id==v)
                    continue
                inherited=inherited_filter(db,model,organization_id,key,v)
                if inherited is None:return None,[f'Not applicable to {model.__tablename__}: {key}']
                q=q.filter(inherited)
            else:q=q.filter(c[key]==v)
    if filters.search:
        strings=[col for col in c if isinstance(col.type,SQLString) and col.name not in HIDDEN_FIELDS][:14]
        if strings:
            pattern=filters.search.replace(chr(92),chr(92)*2).replace('%',chr(92)+'%').replace('_',chr(92)+'_')
            q=q.filter(or_(*(col.ilike('%'+pattern+'%',escape=chr(92)) for col in strings)))
        else:return None,[f'No searchable text in {model.__tablename__}']
    if filters.from_date or filters.to_date:
        name=date_column(model,filters.date_field)
        if not name:return None,[f'No date field in {model.__tablename__}']
        col=c[name]
        if filters.from_date:
            d=filters.from_date
            q=q.filter(col>=datetime.combine(d,time.min,tzinfo=timezone.utc) if isinstance(col.type,SQLDateTime) else col>=d)
        if filters.to_date:
            d=filters.to_date
            q=q.filter(col<datetime.combine(d+timedelta(days=1),time.min,tzinfo=timezone.utc) if isinstance(col.type,SQLDateTime) else col<=d)
    return q,skipped

def enc(v):
    if isinstance(v,Decimal):return float(v)
    if isinstance(v,(date,datetime)):return v.isoformat()
    if isinstance(v,bytes):return '(binary attachment)'
    if isinstance(v,(dict,list,tuple)):return str(v)[:3000]
    return v

def serialize(x,names):return {n:enc(getattr(x,n,None)) for n in names}

def pick_group(model):
    return next((name for name in GROUP_CANDIDATES if name in model.__table__.columns),None)

def primary_key(model):
    return model.__mapper__.primary_key[0]

def count_and_chart(db,model,q):
    col_name=pick_group(model)
    if not col_name:
        return {'type':'bar','title':'Record count','unit':'records','points':[]}
    col=getattr(model,col_name)
    counts=q.with_entities(col.label('label'),func.count(primary_key(model)).label('amount')).group_by(col).order_by(func.count(primary_key(model)).desc()).limit(12).all()
    return {'type':'bar','title':'Records by '+col_name.replace('_',' '),'unit':'records','points':[{'label':str(x[0] or 'Not recorded')[:72],'value':int(x[1])} for x in counts]}

def quantitative_charts(model,q):
    """Additional physical/time-based charts from real aggregated records.

    Monetary measurements are never aggregated without currency grouping.
    """
    cols=model.__table__.columns
    output=[]
    measurement=next((x for x in ('marketable_kg','gross_kg','weight_kg','quantity_kg','head_count','area_ha','quantity') if x in cols),None)
    label_name=pick_group(model)
    if measurement and label_name:
        vals=q.order_by(None).with_entities(cols[label_name],func.sum(cols[measurement])).group_by(cols[label_name]).order_by(func.sum(cols[measurement]).desc()).limit(10).all()
        points=[{'label':str(a or 'Not recorded')[:60],'value':round(float(b or 0),3)} for a,b in vals]
        if points:
            output.append({'type':'bar','title':measurement.replace('_',' ').title()+' by '+label_name.replace('_',' '),
                           'unit':'kg' if measurement.endswith('_kg') else 'ha' if measurement.endswith('_ha') else 'source unit','points':points})
    d=date_column(model)
    if d:
        dt=cols[d]
        year=func.extract('year',dt);month=func.extract('month',dt)
        measure=cols[measurement] if measurement else None
        # Calendar grouping per source date field (not a synthetic report-period timestamp).
        items=(q.order_by(None).with_entities(year,month,func.sum(measure) if measure is not None else func.count(primary_key(model)))
               .filter(dt.isnot(None)).group_by(year,month).order_by(year.asc(),month.asc()).limit(24).all())
        values=[{'label':f'{int(y)}-{int(m):02d}','value':round(float(val or 0),3)} for y,m,val in items if y is not None and m is not None]
        if values:
            output.append({'type':'line','title':('Recorded '+measurement if measurement else 'Record count')+' by month',
                           'unit':'kg' if measurement and measurement.endswith('_kg') else 'ha' if measurement and measurement.endswith('_ha') else 'source unit' if measurement else 'records',
                           'points':values})
    # Financial values in different currencies must NOT share one numerical chart.
    # The finance breakdown records one currency per row instead.
    return output

def metrics(db,model,q):
    names=[c.name for c in model.__table__.columns]
    out={'count':q.order_by(None).count()}
    # Explicit metric unit labeling; never mix unrelated numeric columns.
    for name in VALUE_CANDIDATES:
        if name in names:
            val=q.order_by(None).with_entities(func.sum(getattr(model,name))).scalar()
            out['sum_'+name]=float(val) if val is not None else None
            out['unit_'+name]=('kg' if name.endswith('_kg') else 'ha' if name.endswith('_ha') else 'recorded units (see detail)')
            break
    if 'currency' in names and ('amount' in names or 'gross' in names or 'cost' in names):
        money=next(n for n in ('amount','gross','cost') if n in names)
        sums=q.order_by(None).with_entities(model.currency,func.sum(getattr(model,money))).group_by(model.currency).all()
        out['by_currency']=[{'currency':str(cur or 'Unspecified'),'amount':float(v or 0)} for cur,v in sums]
        out.pop('sum_'+money,None);out.pop('unit_'+money,None)
    if model.__tablename__=='v5_finance_entries':
        sums=q.order_by(None).with_entities(model.currency,model.kind,func.sum(model.amount)).group_by(model.currency,model.kind).all()
        ledger=defaultdict(lambda:dict(cost=0.0,revenue=0.0))
        for cur,kind,amount in sums:
            if kind in ('cost','revenue'):ledger[str(cur or 'Unspecified')][kind]+=float(amount or 0)
        out['ledger']=[dict(currency=cur,**vals,recorded_difference=round(vals['revenue']-vals['cost'],2)) for cur,vals in sorted(ledger.items())]
    return out

def read_report(db,ctx,payload,include_rows=True):
    org=ctx[1]
    sources=report_sources(payload.section_id,ctx)
    if payload.section_id=='custom':raise HTTPException(422,'Choose an individual source from Section Explorer')
    result={'section_id':payload.section_id,'title':DEFINITIONS[payload.section_id][0] if payload.section_id in DEFINITIONS else label(payload.section_id[6:]),
            'generated_at':datetime.now(timezone.utc).isoformat(),'organization':org.name,'filters':payload.filters.model_dump(mode='json'),
            'datasets':[],'omitted':[],'page':payload.page,'page_size':payload.page_size,
            'methodology':['All results are drawn from persisted operational records, not synthetic data.',
                           'Dates on timestamped sources use UTC boundaries unless an organizational timezone integration is configured.',
                           'Numeric totals retain source units; monetary totals are separated by currency.',
                           'A project or other filter excludes datasets that do not carry that relationship rather than showing unfiltered data.']}
    models=models_by_table()
    for tab in sources:
        model=models.get(tab)
        if not model:result['omitted'].append({'table':tab,'reason':'Source ORM model unavailable'});continue
        q,why=q_for(db,model,org.id,payload.filters)
        if q is None:
            result['omitted'].append({'table':tab,'reason':'; '.join(why)});continue
        q=restrict_section(q,model,payload.section_id)
        names=columns(model)
        is_selected=not payload.dataset or payload.dataset==tab
        order_name=payload.sort_by if payload.sort_by in names else ('id' if 'id' in names else names[0])
        sortcol=getattr(model,order_name)
        q_sorted=q.order_by(sortcol.desc() if payload.order=='desc' else sortcol.asc())
        page_rows=[]
        if include_rows and is_selected:
            page_rows=[serialize(r,names) for r in q_sorted.offset((payload.page-1)*payload.page_size).limit(payload.page_size).all()]
        result['datasets'].append({'key':tab,'title':label(tab),'columns':names,'total_rows':q.order_by(None).count(),
                                   'metrics':metrics(db,model,q),'chart':count_and_chart(db,model,q),'additional_charts':quantitative_charts(model,q),
                                   'rows':page_rows,'page':payload.page,'page_size':payload.page_size,'selected':is_selected})
    for omission in result['omitted']:
        omission['title']=label(omission['table'])
    result['analysis_views'],result['analysis_warnings']=analyses(db,ctx,payload,models,set(all_tables(ctx)))
    if include_rows:
        for a in result['analysis_views']:
            a['rows']=a['rows'][:60]
    return result

def export_datasets(db,ctx,payload):
    models=models_by_table();sources=report_sources(payload.section_id,ctx)
    selected=[];omitted=[];total=0
    for tab in sources:
        if payload.dataset and tab!=payload.dataset:continue
        model=models.get(tab)
        if not model:omitted.append({'table':tab,'reason':'ORM source unavailable'});continue
        q,why=q_for(db,model,ctx[1].id,payload.filters)
        if q is None:omitted.append({'table':tab,'reason':'; '.join(why)});continue
        q=restrict_section(q,model,payload.section_id)
        n=q.order_by(None).count();total+=n
        if total>MAX_EXPORT_ROWS:raise HTTPException(413,f'Export contains {total:,}+ rows. Narrow filters to stay within the {MAX_EXPORT_ROWS:,}-row per-export safety limit; no rows have been silently omitted.')
        selected.append((tab,model,q,columns(model),n))
    return selected,omitted,total

def safe_excel(value):
    if value is None:return None
    if isinstance(value,Decimal):return float(value)
    if isinstance(value,datetime) and value.tzinfo is not None:return value.astimezone(timezone.utc).replace(tzinfo=None)
    if isinstance(value,(dict,list)):return str(value)[:5000]
    if isinstance(value,(bytes,bytearray)):return '(binary)'
    if isinstance(value,str):
        if value[:1] in ('=','+','-','@') or value.startswith(('\t','\r','\n')):return "'"+value[:32000]
        return value[:32000]
    return value

def make_excel(report,selected,omitted,mode='complete',charts=True):
    from openpyxl import Workbook
    from openpyxl.chart import BarChart,LineChart,Reference
    from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
    from openpyxl.utils import get_column_letter
    wb=Workbook();summary=wb.active;summary.title='Executive Summary'
    green='174C3B';blue='177EB8';orange='EE963D';purple='8964BB';light='E8F2EE'
    summary.append(['CAPE RANCH | MANAGEMENT INTELLIGENCE'])
    summary.append(['Report',report['title']]);summary.append(['Organization',report['organization']]);summary.append(['Generated (UTC)',report['generated_at']]);
    summary.append(['Scope','Persisted, organization-scoped operational records']);summary.append([])
    for key,val in report['filters'].items():
        if val not in (None,''):summary.append([key.replace('_',' ').title(),str(val)])
    summary.append([]);summary.append(['Reporting dataset','Records','Available','Metric','Currency/Unit'])
    overview_start=summary.max_row+1
    for tab,model,q,names,n in selected:
        ds=next((x for x in report['datasets'] if x['key']==tab),None)
        m=ds['metrics'] if ds else {}
        if m.get('by_currency'):
            for c in m['by_currency']:summary.append([label(tab),n,'Yes',c['amount'],c['currency']])
        else:
            k=next((k for k in m if k.startswith('sum_')),None)
            summary.append([label(tab),n,'Yes',m.get(k) if k else None,m.get('unit_'+k[4:],'records') if k else 'records'])
    overview_end=summary.max_row
    summary.append([]);summary.append(['Excluded sources / limitations'])
    for x in omitted:summary.append([label(x['table']),x['reason']])
    summary.append([]);summary.append(['Methodology'])
    for note in report['methodology']:summary.append([note])
    if mode!='summary':
        used={'Executive Summary'}
        for tab,model,q,names,n in selected:
            sheetname=label(tab)[:29]
            root=sheetname;idx=2
            while sheetname in used:sheetname=(root[:25]+' '+str(idx));idx+=1
            used.add(sheetname);ws=wb.create_sheet(sheetname)
            ws.append([label(tab)+' | '+report['title']]);ws.append(['Data category',label(tab),'Matching records',n]);ws.append([x.replace('_',' ').title() if x!='id' else 'Record ID' for x in names])
            for obj in q.order_by(primary_key(model)).yield_per(1000):
                ws.append([safe_excel(getattr(obj,key,None)) for key in names])
            ws.freeze_panes='A4';ws.auto_filter.ref=f'A3:{get_column_letter(len(names))}{max(3,ws.max_row)}'
            ws.sheet_properties.pageSetUpPr.fitToPage=True
            ws.print_options.horizontalCentered=True
            ws.print_title_rows='1:3'
            for cell in ws[3]:cell.fill=PatternFill('solid',fgColor=green);cell.font=Font(bold=True,color='FFFFFF',size=10);cell.alignment=Alignment(wrap_text=True)
            for colidx,key in enumerate(names,1):ws.column_dimensions[get_column_letter(colidx)].width=min(31,max(13,len(key)+3))
            for cell in ws[1]:cell.font=Font(bold=True,size=14,color=green)
            ws.row_dimensions[1].height=26;ws.row_dimensions[3].height=29
    # Fully linked analytical workbooks: not just raw tables.
    for analysis in report.get('analysis_views',[]):
        if mode=='summary':break
        prefix={'animal_performance':'Animal Growth & Costs','animal_weight_history':'Animal Weight Timeline','feed_usage':'Feed Consumption','inventory_balances':'Stock Balances',
                'crop_yields':'Crop Yield Analysis','machinery_utilization':'Machinery Performance',
                'breeding_followups':'Breeding & Due Dates','livestock_product_yield':'Product Yield',
                'trial_design':'Trial Comparisons','financial_ledger':'Recorded Profit & Loss',
                'project_portfolio':'Project Portfolio','sales_detail':'Sales Analysis',
                'activity_completion':'Routine Completion'}.get(analysis['id'],analysis['title'][:30])
        sheetname=prefix[:31];i=2
        while sheetname in wb.sheetnames:
            sheetname=prefix[:26]+' '+str(i);i+=1
        ws=wb.create_sheet(sheetname)
        ws.append([analysis['title']]);ws.append(['Source-backed derived calculations · see methodology notes'])
        cols=analysis['columns'];ws.append([x.replace('_',' ').title() for x in cols])
        for row in analysis['rows']:
            ws.append([safe_excel(row.get(k)) for k in cols])
        ws.freeze_panes='C4';ws.auto_filter.ref=f'A3:{get_column_letter(len(cols))}{max(3,ws.max_row)}'
        for cell in ws[3]:cell.fill=PatternFill('solid',fgColor=green);cell.font=Font(bold=True,color='FFFFFF');cell.alignment=Alignment(wrap_text=True)
        for i,c in enumerate(cols,1):ws.column_dimensions[get_column_letter(i)].width=min(33,max(13,len(c)+3))
        ws['A1'].font=Font(bold=True,size=14,color=green)
        ws.append([]);ws.append(['Methodology & completeness notes'])
        for note in analysis.get('notes',[]):ws.append([note])
    if report.get('analysis_views'):
        summary.append([]);summary.append(['DETAILED ANALYTICAL TABLES'])
        for a in report['analysis_views']:summary.append([a['title'],a['total_rows'],'Full analysis in dedicated worksheet'])
    if charts and selected:
        chart_ws=wb.create_sheet('Charts & Trends');chart_ws.append(['REPORT CHART DATA | Source record distributions']);chart_ws.append(['Dataset','Category','Number of records'])
        for ds in report['datasets']:
            for chart in [ds['chart'],*(ds.get('additional_charts') or [])]:
                pts=chart.get('points') or []
                if not pts:continue
                first=chart_ws.max_row+1
                for v in pts:chart_ws.append([label(ds['key']),v['label'],v['value']])
                if len(pts)>1:
                    ch=LineChart() if chart.get('type')=='line' else BarChart();ch.style=3+(len(chart_ws._charts)%13);
                    if isinstance(ch,BarChart):ch.type='col';ch.title=ds['title']+' — '+chart['title'];ch.y_axis.title=chart['unit'];ch.x_axis.title=chart['title']
                    ch.add_data(Reference(chart_ws,min_col=3,min_row=first,max_row=chart_ws.max_row))
                    ch.set_categories(Reference(chart_ws,min_col=2,min_row=first,max_row=chart_ws.max_row));ch.width=19;ch.height=9
                    chart_ws.add_chart(ch,'E'+str(2+len(chart_ws._charts)*18))
        for analysis in report.get('analysis_views',[]):
            for graph in analysis.get('charts',[]):
                pts=graph.get('points',[])
                if not pts:continue
                first=chart_ws.max_row+1
                for pt in pts:chart_ws.append([analysis['title'],pt['label'],pt['value']])
                if len(pts)>1:
                    graph_obj=LineChart() if graph['type']=='line' else BarChart()
                    graph_obj.title=analysis['title']+' — '+graph['title']
                    graph_obj.style=4+(len(chart_ws._charts)%10)
                    graph_obj.y_axis.title=graph['unit']
                    graph_obj.add_data(Reference(chart_ws,min_col=3,min_row=first,max_row=chart_ws.max_row))
                    graph_obj.set_categories(Reference(chart_ws,min_col=2,min_row=first,max_row=chart_ws.max_row))
                    graph_obj.width=19;graph_obj.height=9
                    chart_ws.add_chart(graph_obj,'E'+str(2+len(chart_ws._charts)*18))
        for col,width in [('A',34),('B',30),('C',22),('D',4),('E',20)]:chart_ws.column_dimensions[col].width=width
        chart_ws.freeze_panes='A3'
    finance_rows=[]
    for ds in report['datasets']:
        for line in ds.get('metrics',{}).get('ledger',[]):
            finance_rows.append([label(ds['key']),line['currency'],line['revenue'],line['cost'],line['recorded_difference']])
    if finance_rows:
        financial=wb.create_sheet('Financial Breakdown')
        financial.append(['Recorded finance entries by original currency (not audited net income)'])
        financial.append(['Source table','Currency','Recorded revenue','Recorded costs','Recorded difference'])
        for line in finance_rows:
            financial.append(line)
            r=financial.max_row
            financial.cell(r,5,f'=C{r}-D{r}')
        financial.freeze_panes='C3';financial.auto_filter.ref=f'A2:E{financial.max_row}'
        for c in financial[2]:c.fill=PatternFill('solid',fgColor=green);c.font=Font(bold=True,color='FFFFFF')
        for col,w in [('A',34),('B',16),('C',25),('D',25),('E',25)]:financial.column_dimensions[col].width=w
        for row in financial.iter_rows(min_row=3,min_col=3,max_col=5):
            for c in row:c.number_format='#,##0.00;[Red](#,##0.00)'
    details=wb.create_sheet('Filters & Methodology')
    details.append(['Source and methodology notes'])
    details.append(['Organization',report['organization']]);details.append(['Report',report['title']]);details.append(['UTC generated',report['generated_at']]);details.append([])
    for k,v in report['filters'].items():details.append([k,str(v) if v is not None else 'Not set'])
    details.append([]);details.append(['Methodology'])
    for note in report['methodology']:details.append([note])
    details.append([]);details.append(['Excluded datasets'])
    for item in omitted:details.append([label(item['table']),item['reason']])
    details.column_dimensions['A'].width=49;details.column_dimensions['B'].width=87
    details['A1'].font=Font(bold=True,color=green,size=15)
    summary.column_dimensions['A'].width=43;summary.column_dimensions['B'].width=42
    for col in ['C','D','E']:summary.column_dimensions[col].width=23
    summary['A1'].font=Font(bold=True,size=17,color=green)
    for row in [summary[overview_start-1]]:
        for cell in row:cell.fill=PatternFill('solid',fgColor=green);cell.font=Font(bold=True,color='FFFFFF')
    summary.freeze_panes='A2'
    summary.sheet_properties.pageSetUpPr.fitToPage=True
    out=BytesIO();wb.save(out);out.seek(0);return out

def make_pdf(report,selected,omitted,mode='complete',charts=True,paper='a3',orientation='landscape'):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4,A3,letter,landscape,portrait
    from reportlab.lib.styles import getSampleStyleSheet,ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,LongTable,TableStyle,PageBreak,KeepTogether
    from reportlab.graphics.shapes import Drawing
    from reportlab.graphics.charts.barcharts import VerticalBarChart
    from xml.sax.saxutils import escape
    pg={'a4':A4,'a3':A3,'letter':letter}[paper];pg=landscape(pg) if orientation=='landscape' else portrait(pg)
    out=BytesIO();styles=getSampleStyleSheet();body=ParagraphStyle('GHRCell',parent=styles['Normal'],fontSize=7,leading=9,spaceAfter=0)
    small=ParagraphStyle('GHRSmall',parent=styles['Normal'],fontSize=8,leading=11)
    header=ParagraphStyle('GHRHeadCell',parent=body,textColor=colors.white,fontSize=7,leading=9)
    doc=SimpleDocTemplate(out,pagesize=pg,leftMargin=33,rightMargin=33,topMargin=48,bottomMargin=43,title=report['title'])
    story=[];width=pg[0]-66
    def t(v):
        s='—' if v is None else str(enc(v))
        return Paragraph(escape(s[:440]).replace('\n','<br/>'),body)
    def table(title,headers,rows):
        story.append(Paragraph(escape(title),styles['Heading2']))
        if not rows:
            story.append(Paragraph('No matching source records.',small));story.append(Spacer(1,11));return
        # Wide datasets split into readable column blocks; each block retains the ID.
        groups=[];base=['id'] if 'id' in headers else []
        other=[x for x in headers if x not in base]
        chunk=7 if orientation=='landscape' else 4
        for j in range(0,len(other),chunk):groups.append(base+other[j:j+chunk])
        if not groups:groups=[headers]
        for j,part in enumerate(groups):
            if len(groups)>1:story.append(Paragraph('Columns '+str(j+1)+' of '+str(len(groups))+' (same records, matching ID)',small))
            cells=[[Paragraph(escape(x.replace('_',' ').title()),header) for x in part]]
            for row in rows:cells.append([t(row.get(x)) for x in part])
            widths=[width/len(part)]*len(part)
            tbl=LongTable(cells,colWidths=widths,repeatRows=1,hAlign='LEFT')
            tbl.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#174C3B')),
                ('TEXTCOLOR',(0,0),(-1,0),colors.white),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#EEF5F1')]),
                ('LINEBELOW',(0,0),(-1,0),.8,colors.HexColor('#174C3B')),('GRID',(0,1),(-1,-1),.25,colors.HexColor('#CEDDD4')),
                ('VALIGN',(0,0),(-1,-1),'TOP'),('LEFTPADDING',(0,0),(-1,-1),4),('RIGHTPADDING',(0,0),(-1,-1),4),
                ('TOPPADDING',(0,0),(-1,-1),5),('BOTTOMPADDING',(0,0),(-1,-1),5)]))
            story.extend([tbl,Spacer(1,14)])
    story.extend([Paragraph('CAPE RANCH',styles['Title']),Paragraph(escape(report['title']),styles['Heading1']),
                  Paragraph('Comprehensive Management & Intelligence Report',styles['Heading3']),Spacer(1,12),
                  Paragraph('Organization: '+escape(report['organization']),styles['Normal']),
                  Paragraph('Generated (UTC): '+escape(report['generated_at']),styles['Normal']),Spacer(1,14)])
    applied=[f'{k.replace("_"," ")}: {v}' for k,v in report['filters'].items() if v not in (None,'')]
    story.append(Paragraph('Report scope and filters',styles['Heading2']))
    for x in applied or ['All authorized records for this report']:
        story.append(Paragraph(escape(x),small))
    story.append(Spacer(1,11))
    summaryrows=[]
    for table_name,model,q,names,n in selected:
        ds=next((v for v in report['datasets'] if v['key']==table_name),{})
        summaryrows.append({'Source':label(table_name),'Records':n,'Reporting basis':'Persisted records'})
    table('Source record summary',['Source','Records','Reporting basis'],summaryrows)
    for ds in report['datasets']:
        met=ds['metrics'];summary=[{'Measure':'Records','Value':met['count'],'Unit':'records'}]
        for x in met.get('by_currency',[]):summary.append({'Measure':'Recorded amount','Value':x['amount'],'Unit':x['currency']})
        for x in met.get('ledger',[]):summary.extend([{'Measure':k.replace('_',' ').title(),'Value':x[k],'Unit':x['currency']} for k in ('cost','revenue','recorded_difference')])
        for k,v in met.items():
            if k.startswith('sum_'):summary.append({'Measure':k[4:],'Value':v,'Unit':met.get('unit_'+k[4:],'source unit')})
        table(ds['title']+' — summary',['Measure','Value','Unit'],summary)
        for chart in [ds['chart'],*(ds.get('additional_charts') or [])]:
            pts=chart.get('points') or []
            if not charts or len(pts)<2:continue
            maxval=max(float(p['value']) for p in pts)
            minval=min(float(p['value']) for p in pts)
            if maxval<=0 and minval>=0:continue
            # Real recorded group totals, never illustrative demo data.
            d=Drawing(min(width,640),205)
            c=VerticalBarChart();c.x=50;c.y=41;c.width=min(width,640)-100;c.height=122
            c.data=[[float(x['value']) for x in pts[:10]]]
            c.categoryAxis.categoryNames=[x['label'][:12] for x in pts[:10]]
            c.valueAxis.valueMin=min(0,minval*1.1)
            c.valueAxis.valueMax=max(0,maxval*1.16) or 1
            c.barWidth=23;c.categoryAxis.labels.angle=23;c.categoryAxis.labels.boxAnchor='ne'
            shades=['#178E66','#288AC0','#F19A3E','#8D66C6','#D96478','#12AFA9','#9C723F','#5887BF','#AE6E9F','#58A661']
            for i in range(min(len(pts),10)):
                try:c.bars[0,i].fillColor=colors.HexColor(shades[i%len(shades)])
                except (IndexError,TypeError,AttributeError):c.bars[0].fillColor=colors.HexColor(shades[i%len(shades)])
            d.add(c);story.append(d)
            story.append(Paragraph(escape(chart['title']+' — '+chart.get('unit','source units')),small))
            story.append(Spacer(1,12))
    if report.get('analysis_views'):
        story.append(PageBreak())
        story.append(Paragraph('DETAILED MANAGEMENT ANALYSIS',styles['Heading1']))
        story.append(Paragraph('Linked indicators calculated from original operational records',small))
        for analysis in report['analysis_views']:
            # Real time-series graphics for selected individual animals. Table
            # below carries every dated measurement, even if chart is sampled.
            if charts and analysis['id']=='animal_weight_history':
                from reportlab.graphics.shapes import Line, PolyLine, String, Circle
                for graph in analysis.get('charts',[]):
                    values=[p for p in graph.get('points',[]) if p.get('value') is not None]
                    if len(values)<2:continue
                    if len(values)>60:
                        indices=sorted({round(i*(len(values)-1)/59) for i in range(60)})
                        plotted=[values[i] for i in indices]
                    else:plotted=values
                    w=min(width,690);h=180;d=Drawing(w,h)
                    xleft=42;xright=w-15;ybottom=37;ytop=h-18
                    d.add(Line(xleft,ybottom,xleft,ytop,strokeColor=colors.HexColor('#B7CEC2')))
                    d.add(Line(xleft,ybottom,xright,ybottom,strokeColor=colors.HexColor('#B7CEC2')))
                    numbers=[float(p['value']) for p in plotted]
                    low=min(numbers);high=max(numbers);span=high-low or 1
                    xs=[xleft+(xright-xleft)*i/max(1,len(plotted)-1) for i in range(len(plotted))]
                    ys=[ybottom+(numbers[i]-low)/span*(ytop-ybottom) for i in range(len(plotted))]
                    line_color=colors.HexColor('#EE9B42' if 'gain' in graph['title'].lower() else '#258BD3')
                    d.add(PolyLine([v for pair in zip(xs,ys) for v in pair],strokeColor=line_color,strokeWidth=2.5))
                    for x,y in list(zip(xs,ys))[::max(1,len(xs)//15)]:
                        d.add(Circle(x,y,2.5,fillColor=line_color,strokeColor=colors.white,strokeWidth=.6))
                    d.add(String(3,ytop,str(round(high,2)),fontSize=7,fillColor=colors.HexColor('#596E62')))
                    d.add(String(3,ybottom,str(round(low,2)),fontSize=7,fillColor=colors.HexColor('#596E62')))
                    d.add(String(xleft,ybottom-17,str(plotted[0]['label'])[:15],fontSize=8,fillColor=colors.HexColor('#596E62')))
                    d.add(String(xright,ybottom-17,str(plotted[-1]['label'])[:15],fontSize=8,textAnchor='end',fillColor=colors.HexColor('#596E62')))
                    story.append(Paragraph(escape(graph['title']+' ('+graph.get('unit','source unit')+')'),styles['Heading3']))
                    story.append(d)
                    story.append(Spacer(1,10))
            table(analysis['title'],analysis['columns'],analysis['rows'])
            for note in analysis.get('notes',[]):story.append(Paragraph(escape('Methodology: '+note),small))
            story.append(Spacer(1,12))
    if mode!='summary':
        for tab,model,q,names,n in selected:
            # Fully detailed, bounded by explicit guard at export_datasets.
            rows=(serialize(row,names) for row in q.order_by(primary_key(model)).yield_per(500))
            # LongTable requires materialization; explicit upper limit in export_datasets.
            table(label(tab)+' — detailed records ('+str(n)+')',names,list(rows))
    story.append(Paragraph('Limitations and methodology',styles['Heading2']))
    for x in report['methodology']:story.append(Paragraph(escape(x),small))
    for x in omitted:story.append(Paragraph(escape('Excluded '+label(x['table'])+': '+x['reason']),small))
    def canvas_footer(canvas,document):
        canvas.saveState();canvas.setFont('Helvetica',8);canvas.setFillColor(colors.HexColor('#446153'))
        canvas.drawString(33,25,'Cape Ranch  |  '+report['title'][:70]+'  |  Management report')
        canvas.drawRightString(pg[0]-33,25,'Page '+str(document.page));canvas.restoreState()
    doc.build(story,onFirstPage=canvas_footer,onLaterPages=canvas_footer)
    out.seek(0);return out

def make_csv_zip(selected):
    import zipfile
    out=BytesIO()
    with zipfile.ZipFile(out,'w',compression=zipfile.ZIP_DEFLATED) as zipout:
        for tab,model,q,names,n in selected:
            f=StringIO();writer=csv.writer(f);writer.writerow(names)
            for obj in q.order_by(primary_key(model)).yield_per(500):
                writer.writerow([str(enc(safe_excel(getattr(obj,key,None)))) if getattr(obj,key,None) is not None else '' for key in names])
            zipout.writestr(label(tab).replace(' ','_')+'.csv','\ufeff'+f.getvalue())
    out.seek(0);return out
