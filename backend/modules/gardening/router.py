"""Domain-owned operational endpoints.
Uses shared records and permissions without cross-tenant access.
"""
from __future__ import annotations
from datetime import date,datetime
from decimal import Decimal
from fastapi import APIRouter, Depends, HTTPException
from sqlalchemy.orm import Session
from database import get_db
from auth import audit
from modules.organizations.service import context
from modules.crops.models import *
from modules.crops.models import scope,linked,ranch,result,rowdata,saved
from modules.livestock.models import *
from modules.gardening.models import *
from modules.machinery.models import *
router=APIRouter(prefix='/api/v6',tags=['gardening'])


@router.get('/gardening/areas')
def garden_areas(db:Session=Depends(get_db),ctx=Depends(context)):_,o,_=scope(db,ctx,'gardening');return result(db,GardenArea,o)


@router.post('/gardening/areas',status_code=201)
def create_garden(p:GardenIn,db:Session=Depends(get_db),ctx=Depends(context)):
    a,o,_=scope(db,ctx,'gardening',True);ranch(db,o,p.ranch_id);return saved(db,a,GardenArea,GardenArea(organization_id=o.id,**p.model_dump()),'garden_area_created')


@router.get('/gardening/records')
def garden_records(db:Session=Depends(get_db),ctx=Depends(context)):_,o,_=scope(db,ctx,'gardening');return result(db,GardenRecord,o)


@router.post('/gardening/records',status_code=201)
def garden_log(p:GardenLogIn,db:Session=Depends(get_db),ctx=Depends(context)):
    a,o,_=scope(db,ctx,'gardening',True);linked(db,GardenArea,o,p.area_id,'Garden area')
    if db.query(GardenRecord).filter_by(organization_id=o.id,area_id=p.area_id,performed_on=p.performed_on,action=p.action).first():
        raise HTTPException(409,'Activity already recorded for this garden and date. Open the existing record to correct it.')
    return saved(db,a,GardenRecord,GardenRecord(organization_id=o.id,recorded_by=a.id,**p.model_dump()),'garden_activity_recorded')


# V2 APIs are additive; existing /api/v6/gardening/areas and /records stay compatible.
from datetime import timedelta
from sqlalchemy.exc import IntegrityError
from modules.gardening.schemas import (PlantIn, NurseryIn, ScheduleIn, WorkIn,
                                       CompleteIn, HealthIn, ExpenseIn, InspectionIn)
from modules.gardening.service import gardening_dashboard, scheduled_for
from modules.gardening.models import (GardenPlant, GardenNurseryBatch, GardenSchedule,
                                      GardenWork, GardenHealth, GardenExpense, GardenInspection)
from modules.organizations.models import Membership


def _persist(db, actor, obj, action):
    try:
        db.add(obj)
        db.flush()
        audit(db, actor.id, action, obj.__tablename__, obj.id)
        result_data = rowdata(obj)
        db.commit()
    except IntegrityError:
        db.rollback()
        raise HTTPException(409, 'This completion or reference has already been recorded.')
    return {'item': result_data}


def _area(db, org, area_id):
    return linked(db, GardenArea, org, area_id, 'Garden area')


def _plant(db, org, plant_id, area_id):
    if plant_id is None:
        return None
    p = linked(db, GardenPlant, org, plant_id, 'Garden plant')
    if p.area_id != area_id:
        raise HTTPException(422, 'Plant does not belong to the selected garden area')
    return p


def _list(db, org, model, *, limit=200, offset=0, area_id=None):
    q = db.query(model).filter(model.organization_id == org.id)
    if area_id is not None:
        _area(db, org, area_id)
        q = q.filter(model.area_id == area_id)
    return {'total': q.count(), 'items': [rowdata(x) for x in q.order_by(model.id.desc()).offset(offset).limit(limit).all()]}


@router.get('/gardening/v2/dashboard')
def garden_v2_dashboard(db: Session = Depends(get_db), ctx=Depends(context)):
    _, org, _ = scope(db, ctx, 'gardening')
    return gardening_dashboard(db, org.id)


@router.get('/gardening/v2/plants')
def garden_v2_plants(area_id: int | None = None, limit: int = 200, offset: int = 0,
                     db: Session = Depends(get_db), ctx=Depends(context)):
    _, org, _ = scope(db, ctx, 'gardening')
    return _list(db, org, GardenPlant, area_id=area_id, limit=max(1,min(limit,500)), offset=max(offset,0))


@router.post('/gardening/v2/plants', status_code=201)
def garden_v2_create_plant(p: PlantIn, db: Session = Depends(get_db), ctx=Depends(context)):
    actor, org, _ = scope(db, ctx, 'gardening', True)
    _area(db, org, p.area_id)
    return _persist(db, actor, GardenPlant(organization_id=org.id, **p.model_dump()), 'garden_plant_created')


@router.get('/gardening/v2/nursery')
def garden_v2_nursery(area_id: int | None = None, db: Session = Depends(get_db), ctx=Depends(context)):
    _, org, _ = scope(db, ctx, 'gardening')
    return _list(db, org, GardenNurseryBatch, area_id=area_id)


@router.post('/gardening/v2/nursery', status_code=201)
def garden_v2_create_nursery(p: NurseryIn, db: Session = Depends(get_db), ctx=Depends(context)):
    actor, org, _ = scope(db, ctx, 'gardening', True)
    _area(db, org, p.area_id)
    return _persist(db, actor, GardenNurseryBatch(organization_id=org.id, **p.model_dump()), 'garden_nursery_created')


@router.get('/gardening/v2/activities')
def garden_v2_activities(area_id: int | None = None, limit: int = 200, offset: int = 0,
                         db: Session = Depends(get_db), ctx=Depends(context)):
    _, org, _ = scope(db, ctx, 'gardening')
    return _list(db, org, GardenWork, area_id=area_id, limit=max(1,min(limit,500)), offset=max(0,offset))


@router.post('/gardening/v2/activities', status_code=201)
def garden_v2_record_activity(p: WorkIn, db: Session = Depends(get_db), ctx=Depends(context)):
    actor, org, _ = scope(db, ctx, 'gardening', True)
    _area(db, org, p.area_id)
    _plant(db, org, p.plant_id, p.area_id)
    return _persist(db, actor, GardenWork(organization_id=org.id, recorded_by=actor.id,
                                         **p.model_dump()), 'garden_work_completed')


@router.get('/gardening/v2/schedules')
def garden_v2_schedules(area_id: int | None = None, db: Session = Depends(get_db), ctx=Depends(context)):
    _, org, _ = scope(db, ctx, 'gardening')
    return _list(db, org, GardenSchedule, area_id=area_id)


@router.post('/gardening/v2/schedules', status_code=201)
def garden_v2_create_schedule(p: ScheduleIn, db: Session = Depends(get_db), ctx=Depends(context)):
    actor, org, member = scope(db, ctx, 'gardening', True)
    if member.role not in ('org_admin', 'manager'):
        raise HTTPException(403, 'Only garden supervisors can schedule work')
    _area(db, org, p.area_id)
    if p.assigned_to is not None and not db.query(Membership.id).filter_by(
            organization_id=org.id, user_id=p.assigned_to, active=True).first():
        raise HTTPException(422, 'Assignee must be an active member of this organization')
    return _persist(db, actor, GardenSchedule(organization_id=org.id, **p.model_dump()), 'garden_schedule_created')


@router.post('/gardening/v2/schedules/{schedule_id}/complete', status_code=201)
def garden_v2_complete_schedule(schedule_id: int, p: CompleteIn,
                                db: Session = Depends(get_db), ctx=Depends(context)):
    actor, org, member = scope(db, ctx, 'gardening', True)
    task = linked(db, GardenSchedule, org, schedule_id, 'Garden routine')
    if task.assigned_to is not None and task.assigned_to != actor.id and member.role not in ('org_admin', 'manager'):
        raise HTTPException(403, 'Only the assigned gardener or a supervisor can complete this task')
    if not scheduled_for(task, p.performed_on):
        raise HTTPException(422, 'This schedule is not due on the selected date')
    if db.query(GardenWork.id).filter_by(organization_id=org.id, schedule_id=task.id,
                                        performed_on=p.performed_on).first():
        raise HTTPException(409, 'Routine already completed for this date')
    return _persist(db, actor, GardenWork(organization_id=org.id, area_id=task.area_id,
                  schedule_id=task.id, recorded_by=actor.id, action=task.action,
                  shift=task.shift, **p.model_dump()), 'garden_routine_completed')


@router.get('/gardening/v2/health')
def garden_v2_health(area_id: int | None = None, db: Session = Depends(get_db), ctx=Depends(context)):
    _, org, _ = scope(db, ctx, 'gardening')
    return _list(db, org, GardenHealth, area_id=area_id)


@router.post('/gardening/v2/health', status_code=201)
def garden_v2_create_health(p: HealthIn, db: Session = Depends(get_db), ctx=Depends(context)):
    actor, org, _ = scope(db, ctx, 'gardening', True)
    _area(db, org, p.area_id)
    _plant(db, org, p.plant_id, p.area_id)
    return _persist(db, actor, GardenHealth(organization_id=org.id, recorded_by=actor.id,
                                            **p.model_dump()), 'garden_health_recorded')


@router.get('/gardening/v2/expenses')
def garden_v2_expenses(area_id: int | None = None, db: Session = Depends(get_db), ctx=Depends(context)):
    _, org, _ = scope(db, ctx, 'gardening')
    return _list(db, org, GardenExpense, area_id=area_id)


@router.post('/gardening/v2/expenses', status_code=201)
def garden_v2_create_expense(p: ExpenseIn, db: Session = Depends(get_db), ctx=Depends(context)):
    actor, org, _ = scope(db, ctx, 'gardening', True)
    _area(db, org, p.area_id)
    # An inventory item reference is informational; no stock is decremented implicitly.
    from modules.inventory.models import StockItem
    if p.inventory_item_id is not None:
        linked(db, StockItem, org, p.inventory_item_id, 'Inventory item')
    d = p.model_dump()
    d['total_cost'] = p.quantity * p.unit_cost + p.other_cost
    return _persist(db, actor, GardenExpense(organization_id=org.id, recorded_by=actor.id,
                                             **d), 'garden_expense_recorded')


@router.get('/gardening/v2/inspections')
def garden_v2_inspections(area_id: int | None = None, db: Session = Depends(get_db), ctx=Depends(context)):
    _, org, _ = scope(db, ctx, 'gardening')
    return _list(db, org, GardenInspection, area_id=area_id)


@router.post('/gardening/v2/inspections', status_code=201)
def garden_v2_create_inspection(p: InspectionIn, db: Session = Depends(get_db), ctx=Depends(context)):
    actor, org, _ = scope(db, ctx, 'gardening', True)
    _area(db, org, p.area_id)
    d = p.model_dump()
    d['overall_score'] = (p.cleanliness + p.plant_health + p.weed_control + p.irrigation) / 4
    return _persist(db, actor, GardenInspection(organization_id=org.id, inspected_by=actor.id,
                                                **d), 'garden_inspection_recorded')

from modules.gardening.schemas import AreaUpdate, PlantUpdate, NurseryUpdate, HealthUpdate, ScheduleUpdate


def _update(db, actor, obj, changes, action):
    for field, value in changes.items():
        setattr(obj, field, value)
    audit(db, actor.id, action, obj.__tablename__, obj.id)
    result_data = rowdata(obj)
    db.commit()
    return {'item': result_data}


@router.patch('/gardening/v2/areas/{area_id}')
def garden_v2_update_area(area_id: int, p: AreaUpdate, db: Session = Depends(get_db), ctx=Depends(context)):
    actor, org, member = scope(db, ctx, 'gardening', True)
    if member.role not in ('org_admin', 'manager'):
        raise HTTPException(403, 'Only garden supervisors can edit garden areas')
    return _update(db, actor, _area(db, org, area_id), p.model_dump(exclude_unset=True, exclude_none=True), 'garden_area_updated')


@router.patch('/gardening/v2/plants/{plant_id}')
def garden_v2_update_plant(plant_id: int, p: PlantUpdate, db: Session = Depends(get_db), ctx=Depends(context)):
    actor, org, _ = scope(db, ctx, 'gardening', True)
    obj = linked(db, GardenPlant, org, plant_id, 'Garden plant')
    return _update(db, actor, obj, p.model_dump(exclude_unset=True, exclude_none=True), 'garden_plant_updated')


@router.patch('/gardening/v2/nursery/{nursery_id}')
def garden_v2_update_nursery(nursery_id: int, p: NurseryUpdate, db: Session = Depends(get_db), ctx=Depends(context)):
    actor, org, _ = scope(db, ctx, 'gardening', True)
    obj = linked(db, GardenNurseryBatch, org, nursery_id, 'Nursery batch')
    changes = p.model_dump(exclude_unset=True, exclude_none=True)
    germinated = changes.get('germinated', obj.germinated)
    survived = changes.get('survived', obj.survived)
    if not 0 <= survived <= germinated <= obj.seeds_started:
        raise HTTPException(422, 'Survived must be ≤ germinated, and germinated ≤ started')
    return _update(db, actor, obj, changes, 'garden_nursery_updated')


@router.patch('/gardening/v2/health/{health_id}')
def garden_v2_update_health(health_id: int, p: HealthUpdate, db: Session = Depends(get_db), ctx=Depends(context)):
    actor, org, _ = scope(db, ctx, 'gardening', True)
    obj = linked(db, GardenHealth, org, health_id, 'Garden health record')
    return _update(db, actor, obj, p.model_dump(exclude_unset=True, exclude_none=True), 'garden_health_updated')


@router.patch('/gardening/v2/schedules/{schedule_id}')
def garden_v2_update_schedule(schedule_id: int, p: ScheduleUpdate, db: Session = Depends(get_db), ctx=Depends(context)):
    actor, org, member = scope(db, ctx, 'gardening', True)
    if member.role not in ('org_admin', 'manager'):
        raise HTTPException(403, 'Only garden supervisors can manage schedules')
    obj = linked(db, GardenSchedule, org, schedule_id, 'Garden routine')
    return _update(db, actor, obj, p.model_dump(), 'garden_schedule_updated')

@router.get('/gardening/v2/export')
def garden_v2_export(format: str = 'xlsx', db: Session = Depends(get_db), ctx=Depends(context)):
    """Download a tenant-scoped gardening report without persisting server-side files."""
    actor, org, _ = scope(db, ctx, 'gardening')
    if format not in ('xlsx', 'pdf', 'csv'):
        raise HTTPException(422, 'Choose xlsx, pdf or csv')
    from io import BytesIO, StringIO
    import csv
    from fastapi.responses import StreamingResponse
    info = gardening_dashboard(db, org.id)
    sections = [
        ('Garden areas', GardenArea), ('Plant registry', GardenPlant),
        ('Nursery', GardenNurseryBatch), ('Work activities', GardenWork),
        ('Routines', GardenSchedule), ('Health', GardenHealth),
        ('Expenses', GardenExpense), ('Inspections', GardenInspection),
        ('Legacy activities', GardenRecord),
    ]
    data = []
    for title, model in sections:
        q = db.query(model).filter_by(organization_id=org.id)
        data.append((title, q.count(), [rowdata(x) for x in q.order_by(model.id.desc()).limit(500).all()]))
    filename = f'cape-ranch-gardening-{date.today().isoformat()}.{format}'
    headers = {'Content-Disposition': f'attachment; filename="{filename}"', 'Cache-Control': 'no-store'}
    if format == 'csv':
        output = StringIO(newline='')
        writer = csv.writer(output)
        writer.writerow(['Section', 'Records (total)', 'Record ID', 'Field', 'Value'])
        for title, count, records in data:
            if not records:
                writer.writerow([title, count, '', '', ''])
            for r in records:
                for k, v in r.items():
                    # Prevent formula execution when CSV is opened in spreadsheets.
                    safe_value = str(v if v is not None else '')
                    if safe_value.lstrip().startswith(('=', '+', '-', '@')):
                        safe_value = "'" + safe_value
                    writer.writerow([title, count, r.get('id'), k, safe_value])
        return StreamingResponse(iter([output.getvalue().encode('utf-8-sig')]), media_type='text/csv', headers=headers)
    if format == 'xlsx':
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        wb = Workbook()
        sheet = wb.active
        sheet.title = 'Overview'
        sheet.append(['Cape Ranch Gardening Report', date.today().isoformat()])
        for k, v in info['summary'].items():
            sheet.append([k.replace('_', ' ').title(), str(v)])
        for title, count, records in data:
            ws = wb.create_sheet(title[:31])
            ws.append([f'Total records: {count} — most recent 500 included'])
            fields = list(records[0].keys()) if records else ['No records']
            ws.append(fields)
            for row in records:
                ws.append([("'" + str(row[f]) if isinstance(row.get(f), str) and row[f].startswith('=') else row.get(f)) for f in fields])
            for cell in ws[2]:
                cell.font = Font(bold=True)
            ws.freeze_panes = 'A3'
            ws.auto_filter.ref = f'A2:{ws.cell(row=2,column=len(fields)).coordinate}'
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        buf = BytesIO()
        wb.save(buf)
        buf.seek(0)
        return StreamingResponse(buf, media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', headers=headers)
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet
    from xml.sax.saxutils import escape
    buf = BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4)
    sty = getSampleStyleSheet()
    story = [Paragraph('Cape Ranch — Gardening Report', sty['Title']),
             Paragraph(f'Generated {date.today().isoformat()} · {escape(org.name)}', sty['Normal']), Spacer(1, 12)]
    for k, v in info['summary'].items():
        story.append(Paragraph(f'{escape(k.replace("_", " ").title())}: {escape(str(v))}', sty['Normal']))
    story.append(Spacer(1, 14))
    for title, count, rows in data:
        story.append(Paragraph(f'{escape(title)}: {count} records', sty['Heading2']))
        story.append(Paragraph(f'Total {count}; most recent {min(500, len(rows))} included in Excel/CSV export.', sty['Normal']))
    doc.build(story)
    buf.seek(0)
    return StreamingResponse(buf, media_type='application/pdf', headers=headers)
