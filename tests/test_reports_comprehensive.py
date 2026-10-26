"""Real linked farming records -> analytical tables -> PDF/Excel, scoped by organization."""
from __future__ import annotations
import io,os,sys,tempfile
from pathlib import Path
from types import SimpleNamespace
from datetime import datetime,date,timezone
from decimal import Decimal
root=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(root/'backend'))
(root/'tests'/'_test_reports_comprehensive.sqlite3').unlink(missing_ok=True)
os.environ['DATABASE_URL']='sqlite:///'+str(root/'tests'/'_test_reports_comprehensive.sqlite3')
from database import Base,engine,SessionLocal
from modules.reports.v10_service import read_report,export_datasets,make_pdf,make_excel,DEFINITIONS,catalogue
from modules.reports.v10_schemas import ReportRequest,ReportFilter,ExportRequest
from modules.organizations.models import V5Batch,V5Routine
from modules.fattening.models import V5Animal,V5Project,WeightObservation,FeedingLine
from modules.inventory.models import StockItem,StockMovement
from modules.crops.models import V9Project,V9Variety,V9Unit,V9Cycle,V9Harvest,V9Sale,CropField
from modules.machinery.models import Machine,MachineRecord
from modules.finance.models import FinanceEntry
from openpyxl import load_workbook
from pypdf import PdfReader

def test_comprehensive():
 Base.metadata.create_all(engine)
 ctx=(SimpleNamespace(id=1),SimpleNamespace(id=904,name='Example Farm'),SimpleNamespace(role='org_admin',permissions='',_v57_permissions=set()))
 with SessionLocal() as db:
  p=V5Project(organization_id=904,name='Beef Growth',start_date=date(2026,1,1),currency='USD',created_by=1)
  db.add(p);db.flush()
  b=V5Batch(organization_id=904,project_id=p.id,ranch_id=1,name='Cattle Batch',animal_count=0,status='active')
  db.add(b);db.flush()
  animal=V5Animal(organization_id=904,batch_id=b.id,tag='COW-101',species='cattle',breed='Angus',arrival_date=date(2026,1,1),entry_weight_kg=Decimal('250'),current_weight_kg=Decimal('295'),purchase_cost=Decimal('500'),purchase_currency='USD',status='active')
  db.add(animal);db.flush()
  db.add_all([WeightObservation(organization_id=904,animal_id=animal.id,weight_kg=Decimal('270'),recorded_by=1,created_at=datetime(2026,1,21,tzinfo=timezone.utc)),WeightObservation(organization_id=904,animal_id=animal.id,weight_kg=Decimal('295'),recorded_by=1,created_at=datetime(2026,2,15,tzinfo=timezone.utc))])
  item=StockItem(organization_id=904,name='Hay',category='feed',unit='kg',currency='USD',cost_per_unit=Decimal('0.40'),reorder_level=Decimal('100'))
  db.add(item);db.flush()
  routine=V5Routine(organization_id=904,batch_id=b.id,category='feeding',shift='morning',recorded_by=1,client_reference='TEST-RUN-1',created_at=datetime(2026,2,3,tzinfo=timezone.utc))
  db.add(routine);db.flush()
  db.add(FeedingLine(organization_id=904,routine_id=routine.id,item_id=item.id,name='Hay',quantity_kg=Decimal('30')))
  db.add_all([StockMovement(organization_id=904,item_id=item.id,direction='in',quantity=Decimal('500'),unit_cost=Decimal('.40'),reference='SOURCE-1',recorded_by=1),StockMovement(organization_id=904,item_id=item.id,direction='out',quantity=Decimal('30'),unit_cost=Decimal('.40'),reference='ISSUE-1',recorded_by=1)])
  db.add_all([FinanceEntry(organization_id=904,batch_id=b.id,kind='cost',category='feeding',amount=Decimal('75'),currency='USD',created_by=1),FinanceEntry(organization_id=904,batch_id=b.id,kind='revenue',category='sales',amount=Decimal('900'),currency='USD',created_by=1)])
  db.add(V5Animal(organization_id=905,batch_id=b.id,tag='FOREIGN-ANIMAL',species='cattle',purchase_currency='USD',purchase_cost=Decimal(1),status='active'))
  db.commit()
  def report(section,**filters):return read_report(db,ctx,ReportRequest(section_id=section,filters=ReportFilter(**filters)))
  individual=report('individual_animals');performance=next(x for x in individual['analysis_views'] if x['id']=='animal_performance')
  assert len(performance['rows'])==1,performance
  a=performance['rows'][0]
  assert a['tag']=='COW-101' and a['weight_change_kg']==45 and a['latest_weight_kg']==295
  assert a['sale_snapshot_provisional_profit_loss'] is None and a['partial_cost_floor']==500
  one=report('individual_animals',animal_id=animal.id)
  timeline=next(x for x in one['analysis_views'] if x['id']=='animal_weight_history')
  assert len(timeline['rows'])==3, timeline
  assert [x['weight_kg'] for x in timeline['rows']]==[250,270,295]
  assert [x['change_since_previous_kg'] for x in timeline['rows']]==[None,20,25]
  assert round(timeline['rows'][-1]['interval_adg_kg_day'],2)==1.0
  assert timeline['charts'][0]['type']=='line' and len(timeline['charts'][0]['points'])==3
  cat=catalogue(ctx)['sections'];fat=[x for x in cat if x['group']=='Animal Fattening']
  assert {x['id'] for x in fat}=={'fattening_population','individual_animals','fattening_feed','fattening_project_costs'}, fat
  subexport,_,_=export_datasets(db,ctx,ExportRequest(section_id='individual_animals',filters=ReportFilter(animal_id=animal.id)))
  singlepdf=make_pdf(one,subexport,[],mode='complete')
  singletext='\n'.join(p.extract_text() for p in PdfReader(io.BytesIO(singlepdf.getvalue())).pages)
  assert 'Weight Progression' in singletext and 'Liveweight progression over time' in singletext
  singleexcel=make_excel(one,subexport,[],mode='complete')
  singlebook=load_workbook(io.BytesIO(singleexcel.getvalue()))
  assert 'Animal Weight Timeline' in singlebook.sheetnames
  assert len(singlebook['Charts & Trends']._charts)>=1
  feeding=report('feed_consumption');usage=next(x for x in feeding['analysis_views'] if x['id']=='feed_usage')
  assert usage['rows'][0]['quantity_kg']==30
  stock=report('stock_balances');balances=next(x for x in stock['analysis_views'] if x['id']=='inventory_balances')
  assert balances['rows'][0]['movement_balance']==470
  position=report('fattening_project_costs');financial=next(x for x in position['analysis_views'] if x['id']=='project_financial_position')
  assert any(x['recorded_contribution']==825 for x in financial['rows']),financial
  custom=report('table:v5_animals');assert custom['title']=='Individual Fattening Animals'
  exports,omitted,count=export_datasets(db,ctx,ExportRequest(section_id='individual_animals'))
  pdf=make_pdf(individual,exports,omitted,mode='complete');doc=PdfReader(io.BytesIO(pdf.getvalue()));txt='\n'.join(p.extract_text() for p in doc.pages)
  assert 'Individual Animal Fattening' in txt and 'COW-101' in txt and 'Partial cost floor' in txt
  excel=make_excel(individual,exports,omitted,mode='complete');wb=load_workbook(io.BytesIO(excel.getvalue()))
  assert 'Animal Growth & Costs' in wb.sheetnames and 'Executive Summary' in wb.sheetnames and 'Charts & Trends' in wb.sheetnames
  values=str(list(wb['Animal Growth & Costs'].values));assert 'COW-101' in values and 'FOREIGN-ANIMAL' not in values
  assert len(DEFINITIONS)>40
if __name__=='__main__':test_comprehensive();print('Comprehensive farm management reporting passed')
