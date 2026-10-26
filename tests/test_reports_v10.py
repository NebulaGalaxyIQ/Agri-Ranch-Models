"""V10 offline integration + export validation (set DATABASE_URL to a temp SQLite db)."""
import io,os,sys,zipfile
from pathlib import Path
from types import SimpleNamespace
from datetime import datetime,timezone,date
from decimal import Decimal

root=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(root/'backend'))
os.environ['DATABASE_URL']='sqlite:///'+str(root/'tests'/'_test_reports_v10.sqlite3')
from database import Base,engine,SessionLocal
from modules.reports.v10_models import SavedReport,ReportRun
from modules.reports.v10_service import (DEFINITIONS,models_by_table,all_tables,read_report,export_datasets,make_pdf,make_excel,make_csv_zip)
from modules.reports.v10_schemas import ReportRequest,ReportFilter,ExportRequest
from modules.finance.models import FinanceEntry
from sqlalchemy import inspect
from reportlab.pdfbase import pdfmetrics
from openpyxl import load_workbook

def test_v10_end_to_end():
    Base.metadata.create_all(engine)
    org=SimpleNamespace(id=702,name='Report Test Ranch')
    actor=SimpleNamespace(id=7)
    member=SimpleNamespace(role='org_admin',permissions='',_v57_permissions=set())
    ctx=(actor,org,member)
    with SessionLocal() as db:
        db.query(FinanceEntry).filter(FinanceEntry.organization_id.in_([702,703])).delete(synchronize_session=False)
        db.commit()
        db.add_all([
          FinanceEntry(organization_id=702,batch_id=None,animal_id=None,kind='cost',category='feed',amount=Decimal('65'),currency='USD',note='Expenses',created_by=7,created_at=datetime(2026,10,1,tzinfo=timezone.utc)),
          FinanceEntry(organization_id=702,batch_id=None,animal_id=None,kind='revenue',category='sale',amount=Decimal('100'),currency='USD',note='Sample',created_by=7,created_at=datetime(2026,10,4,tzinfo=timezone.utc)),
          FinanceEntry(organization_id=702,batch_id=None,animal_id=None,kind='revenue',category='sale',amount=Decimal('45000'),currency='TZS',note='Other currency',created_by=7,created_at=datetime(2026,10,5,tzinfo=timezone.utc)),
          FinanceEntry(organization_id=703,batch_id=None,animal_id=None,kind='revenue',category='foreign',amount=Decimal('9999999'),currency='USD',note='Must not leak',created_by=7,created_at=datetime(2026,10,5,tzinfo=timezone.utc)),
        ])
        db.commit()
        payload=ReportRequest(section_id='table:v5_finance_entries',filters=ReportFilter(from_date=date(2026,10,1),to_date=date(2026,10,31)))
        result=read_report(db,ctx,payload)
        assert len(result['datasets'])==1,result
        d=result['datasets'][0]
        assert d['total_rows']==3,d
        assert len(d['rows'])==3
        assert {'USD','TZS'}=={x['currency'] for x in d['metrics']['ledger']}
        assert d['metrics']['ledger'][0]['recorded_difference'] not in [9999999]
        scope=ExportRequest(**payload.model_dump())
        exports,omitted,total=export_datasets(db,ctx,scope)
        assert total==3
        pdf=make_pdf(result,exports,omitted)
        content=pdf.getvalue()
        assert content.startswith(b'%PDF'),len(content)
        try:
            from pypdf import PdfReader
        except ImportError:
            assert b'/Type' in content and len(content)>2500
        else:
            pages=PdfReader(io.BytesIO(content)).pages
            assert len(pages)>0
            text=''.join(p.extract_text() for p in pages)
            for word in ['Detailed','100','USD','TZS','feed']:
                assert word.lower() in text.lower(),word
        xls=make_excel(result,exports,omitted)
        book=load_workbook(io.BytesIO(xls.getvalue()),read_only=False,data_only=True)
        assert 'Executive Summary' in book.sheetnames
        assert 'Charts & Trends' in book.sheetnames
        details=[s for s in book.sheetnames if s not in ['Executive Summary','Charts & Trends']][0]
        assert book[details].max_row==6
        allvals=str(list(book[details].values))
        assert '9999999' not in allvals
        assert len(make_csv_zip(exports).getvalue())>10
        # Date boundaries, project filter exclusions and no data handling
        none=read_report(db,ctx,ReportRequest(section_id='finance',filters=ReportFilter(from_date=date(2025,1,1),to_date=date(2025,1,3))))
        assert none['datasets']
        assert all(x['total_rows']==0 for x in none['datasets'])
        grouped=read_report(db,ctx,ReportRequest(section_id='finance',filters=ReportFilter(project_id=1)))
        assert all(x['total_rows']==0 for x in grouped['datasets'])
        assert 'v5_finance_entries' in [x['key'] for x in grouped['datasets']]
        assert len(DEFINITIONS)>=26
        assert len(all_tables(ctx))>40

if __name__=='__main__':
    test_v10_end_to_end();print('V10 reports end-to-end passed')
