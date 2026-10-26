"""End-to-end Gardening V2 API, data, idempotency, permissions and tenant isolation."""
import os
import tempfile
from pathlib import Path

path = Path(tempfile.gettempdir()) / 'cape_gardening_v2_test.sqlite'
if path.exists():
    path.unlink()
os.environ['DATABASE_URL'] = 'sqlite:///' + str(path)
from fastapi.testclient import TestClient
from main import app


def check(response, code=200):
    assert response.status_code == code, (response.status_code, response.text)
    return response.json()


def test_garden_operations_and_isolation():
    with TestClient(app) as c:
        check(c.post('/api/auth/setup', json={
            'full_name':'Platform Owner','email':'gv2owner@example.org','password':'OwnerPassword123!'}))
        organizations=[]
        for n in (1,2):
            r=check(c.post('/api/v5/organizations/register',json={
                'name':f'Garden Org {n}','email':f'gv2org{n}@example.org',
                'admin_name':f'Admin {n}','admin_email':f'gv2admin{n}@example.org',
                'admin_password':'GardenPassword123!'}),201)
            oid=r['organization']['id']
            check(c.put(f'/api/v5/platform/organizations/{oid}/verification',json={'status':'approved'}))
            key=check(c.post('/api/v5/platform/license-keys',json={
                'organization_id':oid,'duration_days':60,'max_employees':25}))['key']
            organizations.append((oid,key))
        check(c.post('/api/auth/logout'))
        check(c.post('/api/auth/login',json={'email':'gv2admin1@example.org','password':'GardenPassword123!'}))
        oid,key=organizations[0];c.headers.update({'X-Organization-ID':str(oid)})
        check(c.post('/api/v5/subscription/redeem',json={'key':key}))
        ranch=check(c.post('/api/v5/ranches',json={'name':'Garden Ranch'}))
        area=check(c.post('/api/v6/gardening/areas',json={'ranch_id':ranch['id'],'name':'Front Flowers','area_m2':100}),201)['item']
        aid=area['id']
        plant=check(c.post('/api/v6/gardening/v2/plants',json={'area_id':aid,'name':'Roses','quantity':30,'condition':'healthy'}),201)['item']
        assert plant['quantity']==30
        check(c.patch(f'/api/v6/gardening/v2/plants/{plant["id"]}',json={'quantity':29,'condition':'stressed'}))
        nursery=check(c.post('/api/v6/gardening/v2/nursery',json={'area_id':aid,'plant_name':'Rose cuttings','seeds_started':50}),201)['item']
        assert c.patch(f'/api/v6/gardening/v2/nursery/{nursery["id"]}',json={'germinated':60}).status_code==422
        check(c.patch(f'/api/v6/gardening/v2/nursery/{nursery["id"]}',json={'germinated':42,'survived':40}))
        morning=check(c.post('/api/v6/gardening/v2/activities',json={
            'area_id':aid,'action':'watering','performed_on':'2026-10-26','shift':'morning',
            'water_litres':85,'direct_cost':2000,'currency':'TZS'}),201)['item']
        evening=check(c.post('/api/v6/gardening/v2/activities',json={
            'area_id':aid,'action':'watering','performed_on':'2026-10-26','shift':'evening',
            'water_litres':75,'direct_cost':2000,'currency':'TZS'}),201)['item']
        assert morning['id'] != evening['id']
        routine=check(c.post('/api/v6/gardening/v2/schedules',json={
            'area_id':aid,'title':'Daily rose watering','action':'watering',
            'frequency':'daily','shift':'afternoon','starts_on':'2026-10-01'}),201)['item']
        sid=routine['id']
        check(c.post(f'/api/v6/gardening/v2/schedules/{sid}/complete',json={
            'performed_on':'2026-10-26','water_litres':25}),201)
        assert c.post(f'/api/v6/gardening/v2/schedules/{sid}/complete',json={'performed_on':'2026-10-26'}).status_code==409
        check(c.post('/api/v6/gardening/v2/health',json={
            'area_id':aid,'plant_id':plant['id'],'issue':'Leaf spot',
            'expense':2000,'currency':'TZS'}),201)
        check(c.post('/api/v6/gardening/v2/expenses',json={
            'area_id':aid,'category':'seed','description':'Rose seeds','quantity':3,
            'unit_cost':400,'other_cost':100,'currency':'USD'}),201)
        check(c.post('/api/v6/gardening/v2/inspections',json={
            'area_id':aid,'cleanliness':5,'plant_health':3,'weed_control':4,'irrigation':4}),201)
        from io import BytesIO
        from openpyxl import load_workbook
        for fmt, ctype in [('xlsx','application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'),('pdf','application/pdf'),('csv','text/csv')]:
            exported=c.get(f'/api/v6/gardening/v2/export?format={fmt}')
            assert exported.status_code==200, exported.text[:250]
            assert ctype in exported.headers['content-type']
            assert len(exported.content)>100
            if fmt=='xlsx':
                wb=load_workbook(BytesIO(exported.content),read_only=True)
                assert {'Overview','Garden areas','Plant registry','Work activities'}<=set(wb.sheetnames)
        dash=check(c.get('/api/v6/gardening/v2/dashboard'))
        assert dash['summary']['areas']==1
        assert dash['summary']['plants']==29
        assert dash['summary']['water_litres']==185
        assert dash['summary']['condition_score']==4
        assert {'currency':'USD','amount':1300} in dash['summary']['expenses_by_currency']
        assert {'currency':'TZS','amount':6000} in dash['summary']['expenses_by_currency']
        assert any(m['month']=='2026-10' and m['activities']==3 for m in dash['monthly'])
        assert c.post('/api/v6/gardening/v2/activities',json={'area_id':999999,'action':'watering'}).status_code==404
        check(c.post('/api/auth/logout'))
        check(c.post('/api/auth/login',json={'email':'gv2admin2@example.org','password':'GardenPassword123!'}))
        c.headers['X-Organization-ID']=str(organizations[1][0])
        check(c.post('/api/v5/subscription/redeem',json={'key':organizations[1][1]}))
        assert check(c.get('/api/v6/gardening/v2/plants'))['items']==[]
        assert check(c.get('/api/v6/gardening/v2/dashboard'))['summary']['areas']==0
        assert c.patch(f'/api/v6/gardening/v2/plants/{plant["id"]}',json={'quantity':1}).status_code==404
        assert c.post('/api/v6/gardening/v2/health',json={'area_id':aid,'issue':'Foreign'}).status_code==404
        assert c.post(f'/api/v6/gardening/v2/schedules/{sid}/complete',json={}).status_code==404
